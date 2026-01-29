import os
from pptx import Presentation
from pptx.util import Inches


def parse_markdown_to_slides(markdown_text):
    slides = []
    # Dummy implementation: split by lines for now.
    lines = markdown_text.split('\n')
    for line in lines:
        if line.startswith('# '):
            title = line[2:]
            slides.append({'type': 'title', 'content': title})
        elif line.startswith('## '):
            content = line[3:]
            slides.append({'type': 'content', 'content': content})
    return slides


def create_title_slide(prs, title):
    slide_layout = prs.slide_layouts[0]  # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    title_placeholder.text = title


def create_content_slide(prs, content):
    slide_layout = prs.slide_layouts[1]  # Content Slide
    slide = prs.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    content_placeholder = slide.shapes.placeholders[1]
    title_placeholder.text = content['content']
    content_placeholder.text = 'Generated content; replace with relevant information.'


def generate_ppt_from_markdown(markdown_text, output_file):
    prs = Presentation()
    slides = parse_markdown_to_slides(markdown_text)
    for slide in slides:
        if slide['type'] == 'title':
            create_title_slide(prs, slide['content'])
        elif slide['type'] == 'content':
            create_content_slide(prs, slide)
    prs.save(output_file)


# Example usage
if __name__ == '__main__':
    markdown_report = '''# My Research Report\n## Introduction\nSome introduction to my research...\n## Conclusion\nFinal thoughts...'''
    generate_ppt_from_markdown(markdown_report, os.path.join(os.getcwd(), 'output.pptx'))